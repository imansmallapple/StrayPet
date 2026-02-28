"""
Build script: assembles all thesis phases into a single DOCX file.

Usage:
    cd thesis/
    python build_thesis.py
"""

import sys, os

# Ensure the thesis directory is on the path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from thesis_utils import setup_document

# Phase imports
from phase1_front_matter import write_front_matter
from phase2_ch1_ch2 import write_ch1_ch2
from phase3_ch3 import write_ch3
from phase4_ch4 import write_ch4
from phase5_ch5_ch6 import write_ch5_ch6
from phase6_appendices import write_appendices


def main():
    print("=" * 60)
    print("Building StrayPet_Thesis.docx ...")
    print("=" * 60)

    doc = Document()
    setup_document(doc)

    print("\n[1/6] Front matter (summary, TOC) ...")
    write_front_matter(doc)

    print("[2/6] Chapter 1 + Chapter 2 ...")
    write_ch1_ch2(doc)

    print("[3/6] Chapter 3 — System design ...")
    write_ch3(doc)

    print("[4/6] Chapter 4 — Implementation ...")
    write_ch4(doc)

    print("[5/6] Chapter 5 + Chapter 6 ...")
    write_ch5_ch6(doc)

    print("[6/6] Appendices + Bibliography ...")
    write_appendices(doc)

    # ── Save ──
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "StrayPet_Thesis.docx")
    doc.save(out_path)

    # ── Statistics ──
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    # Very rough page estimate: ~38 paragraphs per page (A4, 11pt, 1.15 spacing)
    est_pages = max(1, round(para_count / 38))

    print("\n" + "=" * 60)
    print(f"  Saved to: {out_path}")
    print(f"  Paragraphs:      {para_count}")
    print(f"  Tables:           {table_count}")
    print(f"  Estimated pages:  ~{est_pages}")
    print("=" * 60)


if __name__ == "__main__":
    main()
