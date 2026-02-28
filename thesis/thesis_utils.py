"""
Shared formatting utilities for thesis generation.
All helper functions for document styling, paragraphs, tables, etc.
"""

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml


def set_font(run, font_name="Arial", font_size=11, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" '
        f'w:eastAsia="{font_name}" w:cs="{font_name}"/>'
    )
    rPr.insert(0, rFonts)


def add_p(doc, text, font_size=11, bold=False, italic=False,
          alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=4,
          font_name="Arial", color=None, first_line_indent=0.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, font_name, font_size, bold, italic, color)
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)
    return p


def add_p_multi(doc, segments, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=0, space_after=4, first_line_indent=0.5):
    """segments: list of (text, bold, italic) tuples."""
    p = doc.add_paragraph()
    for seg in segments:
        text = seg[0]
        b = seg[1] if len(seg) > 1 else False
        it = seg[2] if len(seg) > 2 else False
        run = p.add_run(text)
        set_font(run, bold=b, italic=it)
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 14, 2: 13, 3: 12}
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_font(run, font_size=sizes.get(level, 11), bold=True, color=(0, 0, 0))
    h.paragraph_format.line_spacing = 1.15
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(8)
    return h


def add_bullet(doc, text, indent=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run("– " + text)
    set_font(run)
    p.paragraph_format.line_spacing = 1.15
    if indent > 0:
        p.paragraph_format.left_indent = Cm(indent * 1.27)
    return p


def add_code_block(doc, code_text, caption=""):
    if caption:
        add_p(doc, caption, font_size=9, italic=True, first_line_indent=0, space_before=8)
    for line in code_text.split('\n'):
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_font(run, font_name="Consolas", font_size=9)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(1)


def add_table(doc, title_num, title_text, headers, rows):
    add_p(doc, f"Table {title_num}. {title_text}", font_size=9, bold=False,
          first_line_indent=0, space_before=12, space_after=4)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_font(run, font_size=9, bold=True)
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E2F3"/>')
        cell._tc.get_or_add_tcPr().append(shading)
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(cell_text))
            set_font(run, font_size=9)
    return table


def add_figure(doc, fig_num, caption, instruction=""):
    add_p(doc, f"[INSERT FIGURE: {instruction or caption}]",
          font_size=10, italic=True, color=(128, 128, 128),
          alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=0,
          space_before=12, space_after=4)
    add_p(doc, f"Figure {fig_num}. {caption}", font_size=9,
          first_line_indent=0, space_before=2, space_after=8,
          alignment=WD_ALIGN_PARAGRAPH.LEFT)


def add_source(doc, source_text):
    add_p(doc, f"Source: {source_text}", font_size=9, italic=True,
          first_line_indent=0, space_before=0, space_after=8)


def page_break(doc):
    doc.add_page_break()


def setup_document(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.15
    rPr = style.element.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="Arial" w:hAnsi="Arial" '
        f'w:eastAsia="Arial" w:cs="Arial"/>'
    )
    rPr.insert(0, rFonts)

    for level, size in [(1, 14), (2, 13), (3, 12)]:
        sn = f'Heading {level}'
        if sn in doc.styles:
            hs = doc.styles[sn]
            hs.font.name = 'Arial'
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = RGBColor(0, 0, 0)
            hs.paragraph_format.line_spacing = 1.15

    section = doc.sections[0]
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)
    section.page_width = Mm(210)
    section.page_height = Mm(297)

    sectPr = section._sectPr
    mirror = parse_xml(f'<w:mirrorMargins {nsdecls("w")}/>')
    sectPr.append(mirror)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run1 = fp.add_run()
    run1._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = fp.add_run()
    run2._element.append(parse_xml(
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'
    ))
    run3 = fp.add_run()
    run3._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    for r in [run1, run2, run3]:
        r.font.name = "Arial"
        r.font.size = Pt(9)
